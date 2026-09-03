"""Spreadsheet import/export service using openpyxl."""
from __future__ import annotations

import io
import json
import os
import re
import uuid
from datetime import date, datetime
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string


def parse_excel_to_spreadsheet_json(file_stream_or_path: Any) -> str:
    """Parse an uploaded Excel/CSV file into standard EDMS spreadsheet JSON format."""
    try:
        if isinstance(file_stream_or_path, str) and file_stream_or_path.lower().endswith(".csv"):
            import csv
            sheets = []
            sheet_data = {
                "id": "sheet_1",
                "name": "Sheet1",
                "rowCount": 100,
                "colCount": 26,
                "cells": {},
                "styles": {},
                "images": [],
                "mergedCells": [],
                "columnWidths": {},
                "rowHeights": {}
            }
            with open(file_stream_or_path, "r", encoding="utf-8-sig", errors="replace") as f:
                reader = csv.reader(f)
                r_idx = 0
                max_cols = 26
                for row in reader:
                    if len(row) > max_cols:
                        max_cols = len(row)
                    for c_idx, val in enumerate(row):
                        if val != "":
                            parsed_val = val
                            try:
                                if "." in val:
                                    parsed_val = float(val)
                                else:
                                    parsed_val = int(val)
                            except ValueError:
                                parsed_val = val
                            sheet_data["cells"][f"{r_idx}_{c_idx}"] = {
                                "v": parsed_val,
                                "m": str(parsed_val),
                                "t": "n" if isinstance(parsed_val, (int, float)) else "s"
                            }
                    r_idx += 1
                sheet_data["rowCount"] = max(r_idx + 20, 50)
                sheet_data["colCount"] = max(max_cols + 5, 26)
            sheets.append(sheet_data)
            return json.dumps({"type": "spreadsheet", "activeSheetId": "sheet_1", "sheets": sheets})

        # Load with openpyxl
        wb = openpyxl.load_workbook(file_stream_or_path, data_only=False)
        sheets: List[Dict[str, Any]] = []

        for s_idx, sheetname in enumerate(wb.sheetnames):
            ws = wb[sheetname]
            sheet_id = f"sheet_{s_idx + 1}"
            
            cells: Dict[str, Any] = {}
            styles: Dict[str, Any] = {}
            merged_cells: List[Dict[str, int]] = []
            column_widths: Dict[str, int] = {}
            row_heights: Dict[str, int] = {}
            images: List[Dict[str, Any]] = []

            # Merged cell ranges
            for rng in ws.merged_cells.ranges:
                merged_cells.append({
                    "r": rng.min_row - 1,
                    "c": rng.min_col - 1,
                    "rs": rng.max_row - rng.min_row + 1,
                    "cs": rng.max_col - rng.min_col + 1
                })

            # Column dimensions
            for col_letter, col_dim in ws.column_dimensions.items():
                if col_dim.width:
                    try:
                        col_idx = column_index_from_string(col_letter) - 1
                        column_widths[str(col_idx)] = int(col_dim.width * 8)
                    except:
                        pass

            # Row dimensions
            for row_idx, row_dim in ws.row_dimensions.items():
                if row_dim.height:
                    row_heights[str(row_idx - 1)] = int(row_dim.height * 1.33)

            # Images in worksheet
            try:
                if hasattr(ws, "_images"):
                    for img in ws._images:
                        img_dict = {
                            "id": f"img_{uuid.uuid4().hex[:8]}",
                            "r": getattr(img.anchor, "_from", None).row if hasattr(img.anchor, "_from") else 0,
                            "c": getattr(img.anchor, "_from", None).col if hasattr(img.anchor, "_from") else 0,
                            "width": int(getattr(img, "width", 200)),
                            "height": int(getattr(img, "height", 150)),
                        }
                        if hasattr(img, "_data"):
                            import base64
                            b64 = base64.b64encode(img._data()).decode("utf-8")
                            img_dict["src"] = f"data:image/png;base64,{b64}"
                            images.append(img_dict)
            except Exception as img_err:
                print(f"[Spreadsheet Parser] Image extraction notice: {img_err}")

            # Iterate rows and cells
            max_r = max(ws.max_row or 1, 50)
            max_c = max(ws.max_column or 1, 26)

            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    if cell.value is not None or cell.comment:
                        r = cell.row - 1
                        c = cell.column - 1
                        cell_key = f"{r}_{c}"
                        
                        val = cell.value
                        val_type = "s"
                        formula = None
                        
                        if isinstance(val, str) and val.startswith("="):
                            formula = val
                            val_type = "f"
                        elif isinstance(val, (int, float)):
                            val_type = "n"
                        elif isinstance(val, bool):
                            val_type = "b"
                        elif isinstance(val, (date, datetime)):
                            val = val.isoformat()
                            val_type = "d"

                        cell_obj: Dict[str, Any] = {
                            "v": val,
                            "m": str(val) if val is not None else "",
                            "t": val_type
                        }
                        if formula:
                            cell_obj["f"] = formula

                        # Styling extraction
                        style_obj: Dict[str, Any] = {}
                        if cell.font:
                            if cell.font.bold: style_obj["bl"] = 1
                            if cell.font.italic: style_obj["it"] = 1
                            if cell.font.underline: style_obj["un"] = 1
                            if cell.font.strike: style_obj["cl"] = 1
                            if cell.font.size: style_obj["fs"] = int(cell.font.size)
                            if cell.font.color and hasattr(cell.font.color, "rgb") and isinstance(cell.font.color.rgb, str):
                                rgb = cell.font.color.rgb
                                if len(rgb) == 8:
                                    style_obj["fc"] = f"#{rgb[2:]}"
                                elif len(rgb) == 6:
                                    style_obj["fc"] = f"#{rgb}"

                        if cell.fill and hasattr(cell.fill, "start_color") and cell.fill.start_color:
                            sc = cell.fill.start_color
                            if hasattr(sc, "rgb") and isinstance(sc.rgb, str) and sc.rgb != "00000000":
                                rgb = sc.rgb
                                if len(rgb) == 8:
                                    style_obj["bg"] = f"#{rgb[2:]}"
                                elif len(rgb) == 6:
                                    style_obj["bg"] = f"#{rgb}"

                        if cell.alignment:
                            if cell.alignment.horizontal:
                                style_obj["ht"] = cell.alignment.horizontal
                            if cell.alignment.vertical:
                                style_obj["vt"] = cell.alignment.vertical
                            if cell.alignment.wrap_text:
                                style_obj["tb"] = 1

                        if style_obj:
                            styles[cell_key] = style_obj

                        cells[cell_key] = cell_obj

            sheets.append({
                "id": sheet_id,
                "name": sheetname,
                "rowCount": max(max_r + 20, 60),
                "colCount": max(max_c + 5, 26),
                "cells": cells,
                "styles": styles,
                "images": images,
                "mergedCells": merged_cells,
                "columnWidths": column_widths,
                "rowHeights": row_heights
            })

        if not sheets:
            sheets.append({
                "id": "sheet_1",
                "name": "Sheet1",
                "rowCount": 60,
                "colCount": 26,
                "cells": {},
                "styles": {},
                "images": [],
                "mergedCells": [],
                "columnWidths": {},
                "rowHeights": {}
            })

        return json.dumps({
            "type": "spreadsheet",
            "activeSheetId": sheets[0]["id"] if sheets else "sheet_1",
            "sheets": sheets
        })

    except Exception as e:
        print(f"[Spreadsheet Parser] Error: {e}")
        return json.dumps({
            "type": "spreadsheet",
            "activeSheetId": "sheet_1",
            "sheets": [{
                "id": "sheet_1",
                "name": "Sheet1",
                "rowCount": 60,
                "colCount": 26,
                "cells": {},
                "styles": {},
                "images": [],
                "mergedCells": [],
                "columnWidths": {},
                "rowHeights": {}
            }]
        })


def export_spreadsheet_json_to_excel_bytes(content_json_str: str) -> bytes:
    """Convert EDMS spreadsheet JSON into binary .xlsx bytes."""
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    try:
        data = json.loads(content_json_str) if isinstance(content_json_str, str) else (content_json_str or {})
        sheets_data = data.get("sheets", [])
        if not sheets_data:
            ws = wb.create_sheet(title="Sheet1")
            ws.cell(row=1, column=1, value="")
        else:
            for s in sheets_data:
                sheet_name = (s.get("name") or "Sheet").strip()[:31]
                sheet_name = re.sub(r'[\\/*?:\[\]]', '', sheet_name) or "Sheet"
                ws = wb.create_sheet(title=sheet_name)

                cells = s.get("cells", {})
                styles = s.get("styles", {})
                merged_cells = s.get("mergedCells", [])
                col_widths = s.get("columnWidths", {})
                row_heights = s.get("rowHeights", {})

                # Column widths
                for col_str, width_px in col_widths.items():
                    try:
                        col_idx = int(col_str) + 1
                        col_letter = get_column_letter(col_idx)
                        ws.column_dimensions[col_letter].width = max(int(width_px) / 8, 8)
                    except:
                        pass

                # Row heights
                for row_str, height_px in row_heights.items():
                    try:
                        row_idx = int(row_str) + 1
                        ws.row_dimensions[row_idx].height = max(int(height_px) / 1.33, 15)
                    except:
                        pass

                # Merged cells
                for m in merged_cells:
                    try:
                        r = int(m.get("r", 0)) + 1
                        c = int(m.get("c", 0)) + 1
                        rs = int(m.get("rs", 1))
                        cs = int(m.get("cs", 1))
                        if rs > 1 or cs > 1:
                            ws.merge_cells(start_row=r, start_column=c, end_row=r + rs - 1, end_column=c + cs - 1)
                    except Exception as merge_err:
                        print(f"[Spreadsheet Export] Merge cell error: {merge_err}")

                # Cell values and styles
                for cell_key, cell_info in cells.items():
                    try:
                        r_str, c_str = cell_key.split("_")
                        r = int(r_str) + 1
                        c = int(c_str) + 1
                        
                        target_cell = ws.cell(row=r, column=c)
                        
                        formula = cell_info.get("f")
                        val = cell_info.get("v")
                        if formula:
                            target_cell.value = formula if formula.startswith("=") else f"={formula}"
                        else:
                            target_cell.value = val

                        st = styles.get(cell_key, {})
                        if st:
                            font_kwargs: Dict[str, Any] = {}
                            if st.get("bl"): font_kwargs["bold"] = True
                            if st.get("it"): font_kwargs["italic"] = True
                            if st.get("un"): font_kwargs["underline"] = "single"
                            if st.get("cl"): font_kwargs["strike"] = True
                            if st.get("fs"): font_kwargs["size"] = int(st.get("fs"))
                            if st.get("fc"):
                                hex_col = st["fc"].lstrip("#")
                                if len(hex_col) == 6:
                                    font_kwargs["color"] = hex_col
                            if font_kwargs:
                                target_cell.font = Font(**font_kwargs)

                            if st.get("bg"):
                                bg_hex = st["bg"].lstrip("#")
                                if len(bg_hex) == 6:
                                    target_cell.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")

                            align_kwargs: Dict[str, Any] = {}
                            if st.get("ht"): align_kwargs["horizontal"] = st["ht"]
                            if st.get("vt"): align_kwargs["vertical"] = st["vt"]
                            if st.get("tb"): align_kwargs["wrap_text"] = True
                            if align_kwargs:
                                target_cell.alignment = Alignment(**align_kwargs)

                    except Exception as cell_err:
                        print(f"[Spreadsheet Export] Cell write error at {cell_key}: {cell_err}")

    except Exception as e:
        print(f"[Spreadsheet Export] Workbook generation error: {e}")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
